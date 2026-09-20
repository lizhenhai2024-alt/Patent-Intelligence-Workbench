"""CSV and Excel exports for local patent library views."""

from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from app.library.models import LibraryPatent

EXPORT_COLUMNS = (
    "publication_number",
    "jurisdiction",
    "kind_code",
    "family_key",
    "family_type",
    "source_family_id",
    "title",
    "application_number",
    "grant_number",
    "filing_date",
    "publication_date",
    "grant_date",
    "language",
    "earliest_priority_number",
    "earliest_priority_date",
    "priorities",
    "classifications",
    "original_assignees",
    "current_assignees",
    "company_groups",
    "technology_topics",
    "projects",
    "tags",
    "pdf_paths",
    "pdf_providers",
    "watch_rule_ids",
    "provenance",
    "favorite",
    "note",
    "first_seen_at",
    "last_seen_at",
    "source",
)


def patent_to_export_row(patent: LibraryPatent) -> dict[str, str | bool]:
    return {
        "publication_number": patent.publication_number,
        "jurisdiction": patent.jurisdiction,
        "kind_code": patent.kind_code or "",
        "family_key": patent.family_key or "",
        "family_type": patent.family_type or "",
        "source_family_id": patent.source_family_id or "",
        "title": patent.title or "",
        "application_number": patent.application_number or "",
        "grant_number": patent.grant_number or "",
        "filing_date": patent.filing_date.isoformat() if patent.filing_date else "",
        "publication_date": (
            patent.publication_date.isoformat() if patent.publication_date else ""
        ),
        "grant_date": patent.grant_date.isoformat() if patent.grant_date else "",
        "language": patent.language or "",
        "earliest_priority_number": patent.earliest_priority_number or "",
        "earliest_priority_date": (
            patent.earliest_priority_date.isoformat()
            if patent.earliest_priority_date
            else ""
        ),
        "priorities": " | ".join(
            _priority_text(priority) for priority in patent.priorities
        ),
        "classifications": " | ".join(
            f"{item.system}:{item.code}{'*' if item.is_main else ''}"
            for item in patent.classifications
        ),
        "original_assignees": " | ".join(patent.original_assignees),
        "current_assignees": " | ".join(patent.current_assignees),
        "company_groups": " | ".join(patent.company_groups),
        "technology_topics": " | ".join(patent.technology_topics),
        "projects": " | ".join(patent.projects),
        "tags": " | ".join(patent.tags),
        "pdf_paths": " | ".join(str(document.path) for document in patent.documents),
        "pdf_providers": " | ".join(
            document.provider or "" for document in patent.documents
        ),
        "watch_rule_ids": " | ".join(patent.watch_rule_ids),
        "provenance": " | ".join(
            f"{source.source_type}:{source.source_ref}" for source in patent.provenance
        ),
        "favorite": patent.favorite,
        "note": patent.note or "",
        "first_seen_at": patent.first_seen_at.isoformat(),
        "last_seen_at": patent.last_seen_at.isoformat(),
        "source": patent.source or "",
    }


def export_csv(
    patents: tuple[LibraryPatent, ...],
    destination: str | Path,
) -> Path:
    path = Path(destination)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=EXPORT_COLUMNS)
        writer.writeheader()
        for patent in patents:
            writer.writerow(patent_to_export_row(patent))
    return path


def export_xlsx(
    patents: tuple[LibraryPatent, ...],
    destination: str | Path,
) -> Path:
    path = Path(destination)
    path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Patent Library"

    for column_index, name in enumerate(EXPORT_COLUMNS, start=1):
        cell = sheet.cell(row=1, column=column_index, value=name)
        cell.font = Font(bold=True)

    rows = tuple(patent_to_export_row(patent) for patent in patents)
    for row_index, values in enumerate(rows, start=2):
        for column_index, name in enumerate(EXPORT_COLUMNS, start=1):
            sheet.cell(
                row=row_index,
                column=column_index,
                value=values[name],
            )

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    for column_index, name in enumerate(EXPORT_COLUMNS, start=1):
        values = [name]
        values.extend(str(row[name]) for row in rows[:200])
        max_length = min(max((len(value) for value in values), default=10) + 2, 60)
        sheet.column_dimensions[get_column_letter(column_index)].width = max_length

    workbook.save(path)
    return path


def _priority_text(priority) -> str:
    date_text = priority.priority_date.isoformat() if priority.priority_date else ""
    parts = [priority.number, priority.country, date_text, priority.priority_type or ""]
    return "/".join(part for part in parts if part)

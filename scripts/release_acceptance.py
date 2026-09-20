"""Offline V1 release acceptance.

Runs the core product chain without external services:
Search -> Family -> Download -> Patent Watch -> Local Library -> CSV/XLSX.
"""

from __future__ import annotations

import argparse
import asyncio
import json
from datetime import UTC, date, datetime
from pathlib import Path

from openpyxl import load_workbook

from app.core.company_registry import CompanyRegistry
from app.core.patent_number import normalize_patent_number
from app.domain.family import (
    Classification,
    FamilyType,
    PatentFamily,
    PatentPublication,
    PriorityClaim,
)
from app.domain.search import SearchHit, SearchPage
from app.downloads.base import PdfPayload
from app.downloads.family import FamilyDownloader
from app.downloads.manager import DownloadManager
from app.library.models import LibraryQuery
from app.library.service import PatentLibraryService
from app.library.store import SQLitePatentLibrary
from app.providers.base import ProviderCapability, ProviderInfo
from app.services.family_resolver import FamilyResolver
from app.services.search_service import SearchService
from app.watch.engine import PatentWatchEngine
from app.watch.models import WatchEventType, WatchRule
from app.watch.state import SQLiteWatchStateStore


class AcceptanceSearchProvider:
    info = ProviderInfo(
        name="ACCEPTANCE_SEARCH",
        capabilities=frozenset(
            {
                ProviderCapability.SEARCH,
                ProviderCapability.PUBLICATION_LOOKUP,
            }
        ),
    )

    def __init__(self) -> None:
        self.hits: tuple[SearchHit, ...] = ()

    async def lookup_publication(self, publication):
        return SearchPage(
            hits=(
                SearchHit(
                    publication_number=publication.canonical,
                    jurisdiction=publication.jurisdiction,
                    kind_code=publication.kind_code,
                    source=self.info.name,
                ),
            ),
            total_result_count=1,
            range_begin=1,
            range_end=1,
        )

    async def search_publications(
        self,
        expression,
        *,
        page_size=25,
        page_start=1,
    ):
        del expression
        page = self.hits[page_start - 1 : page_start - 1 + page_size]
        end = page_start + len(page) - 1 if page else None
        return SearchPage(
            hits=tuple(page),
            total_result_count=len(self.hits),
            range_begin=page_start if page else None,
            range_end=end,
        )


class AcceptanceFamilyProvider:
    info = ProviderInfo(
        name="ACCEPTANCE_FAMILY",
        capabilities=frozenset({ProviderCapability.FAMILY_SIMPLE}),
    )

    def __init__(self) -> None:
        self.families: dict[str, PatentFamily] = {}

    async def get_family(self, publication, family_type):
        assert family_type is FamilyType.DOCDB_SIMPLE
        return self.families[publication.canonical]


class AcceptancePdfProvider:
    name = "ACCEPTANCE_PDF"

    def supports(self, publication) -> bool:
        return True

    async def fetch_pdf(self, publication):
        return PdfPayload(
            provider=self.name,
            source_url=f"https://acceptance.invalid/{publication.canonical}.pdf",
            data=(
                b"%PDF-1.7\n"
                + publication.canonical.encode("ascii")
                + b"\n%%EOF"
            ),
        )


def _registry() -> CompanyRegistry:
    return CompanyRegistry.from_dict(
        {
            "companies": [
                {
                    "group_id": "testco",
                    "display_name": "Test Company",
                    "core_watch": True,
                    "entities": [
                        {
                            "name": "Test Company Ltd.",
                            "relation": "SAME_ENTITY",
                        }
                    ],
                }
            ]
        }
    )


def _family(*members: str) -> PatentFamily:
    priority = PriorityClaim(
        number="JP2022000456",
        country="JP",
        priority_date=date(2022, 3, 4),
        priority_type="national",
    )
    publications = []
    for number in members:
        normalized = normalize_patent_number(number)
        publications.append(
            PatentPublication(
                publication_number=normalized.canonical,
                jurisdiction=normalized.jurisdiction,
                kind_code=normalized.kind_code,
                title="Electronically controlled damper",
                publication_date=date(2026, 9, 1),
                original_assignees=("Test Company Ltd.",),
                priorities=(priority,),
                classifications=(
                    Classification("IPC", "F16F9/46", True),
                ),
            )
        )
    return PatentFamily(
        family_type=FamilyType.DOCDB_SIMPLE,
        source="ACCEPTANCE_FAMILY",
        source_family_id="ACCEPTANCE-F1",
        members=publications,
    )


async def run_acceptance(root: Path) -> dict[str, object]:
    root.mkdir(parents=True, exist_ok=True)
    db_path = root / "workbench.db"
    library = SQLitePatentLibrary(db_path)
    watch_store = SQLiteWatchStateStore(db_path)

    try:
        library_service = PatentLibraryService(library)
        search_provider = AcceptanceSearchProvider()
        family_provider = AcceptanceFamilyProvider()
        search_service = SearchService(
            provider=search_provider,
            company_registry=_registry(),
        )
        resolver = FamilyResolver([family_provider])

        jp = "JP2024000123A"
        wo = "WO2024123456A1"
        us = "US20240123456A1"

        search_provider.hits = (
            SearchHit(
                publication_number=jp,
                jurisdiction="JP",
                kind_code="A",
                title="Electronically controlled damper",
                applicants=("Test Company Ltd.",),
                publication_date=date(2026, 9, 1),
                source=search_provider.info.name,
            ),
        )
        search_response = await search_service.search(
            "pilot valve",
            company="Test Company",
            technology_terms=("pilot valve",),
            jurisdictions=("JP", "WO", "US"),
            page_size=100,
        )
        library_service.ingest_search_response(
            search_response,
            technology_topics=("pilot_control_valve",),
            projects=("RC Acceptance",),
            tags=("acceptance",),
            seen_at=datetime(2026, 9, 20, 8, 0, tzinfo=UTC),
        )

        initial_family = _family(jp, wo)
        family_provider.families[jp] = initial_family
        resolution = await resolver.resolve(
            normalize_patent_number(jp),
            FamilyType.DOCDB_SIMPLE,
        )
        library_service.ingest_family(
            resolution.family,
            company_groups=("testco",),
            technology_topics=("pilot_control_valve",),
            projects=("RC Acceptance",),
            seen_at=datetime(2026, 9, 20, 8, 5, tzinfo=UTC),
        )

        downloader = FamilyDownloader(
            DownloadManager([AcceptancePdfProvider()])
        )
        download_summary = await downloader.download_family(
            resolution.family,
            root / "downloads",
        )
        library_service.ingest_download_summary(
            download_summary,
            seen_at=datetime(2026, 9, 20, 8, 10, tzinfo=UTC),
        )
        assert download_summary.succeeded == 2

        def archive(rule, event) -> None:
            library_service.ingest_watch_event(
                event,
                company_group=rule.company_group,
                technology_topics=rule.technology_terms,
            )

        watch_engine = PatentWatchEngine(
            search_service=search_service,
            family_resolver=resolver,
            state_store=watch_store,
            event_sink=archive,
        )
        rule = WatchRule(
            rule_id="acceptance-watch",
            name="Acceptance watch",
            company_group="testco",
            technology_terms=("pilot_control_valve",),
        )

        # Baseline: no historical notification.
        search_provider.hits = (
            SearchHit(
                publication_number=jp,
                jurisdiction="JP",
                title="Electronically controlled damper",
                publication_date=date(2026, 9, 1),
            ),
        )
        family_provider.families[jp] = initial_family
        baseline = await watch_engine.run_rule(
            rule,
            now=datetime(2026, 9, 20, 9, 0, tzinfo=UTC),
        )
        assert baseline.events == ()

        expanded_family = _family(jp, wo, us)
        family_provider.families[us] = expanded_family
        search_provider.hits = (
            SearchHit(
                publication_number=us,
                jurisdiction="US",
                kind_code="A1",
                title="Electronically controlled damper",
                publication_date=date(2026, 9, 21),
            ),
        )
        incremental = await watch_engine.run_rule(
            rule,
            now=datetime(2026, 9, 21, 9, 0, tzinfo=UTC),
        )
        assert len(incremental.events) == 1
        assert (
            incremental.events[0].event_type
            is WatchEventType.NEW_FAMILY_MEMBER
        )
        assert incremental.events[0].publication_number == us

        jp_record = library.get_patent(jp)
        us_record = library.get_patent(us)
        assert jp_record is not None
        assert us_record is not None
        assert jp_record.family_key is not None
        assert len(jp_record.documents) == 1
        assert us_record.watch_rule_ids == ("acceptance-watch",)

        pdf_records = library.query(
            LibraryQuery(has_pdf=True, limit=100)
        )
        assert {item.publication_number for item in pdf_records} == {jp, wo}

        csv_path = library_service.export(
            root / "exports" / "acceptance.csv",
            query=LibraryQuery(limit=100),
        )
        xlsx_path = library_service.export(
            root / "exports" / "acceptance.xlsx",
            query=LibraryQuery(limit=100),
        )
        assert csv_path.is_file()
        assert xlsx_path.is_file()

        workbook = load_workbook(xlsx_path, read_only=True)
        try:
            sheet = workbook["Patent Library"]
            assert sheet.max_row >= 4
        finally:
            workbook.close()

        return {
            "database": str(db_path),
            "patents": library.count_patents(),
            "family_download_success": download_summary.succeeded,
            "watch_events": len(incremental.events),
            "csv": str(csv_path),
            "xlsx": str(xlsx_path),
        }
    finally:
        library.close()
        watch_store.close()


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--data-dir",
        required=True,
        help="Temporary directory used for offline release acceptance.",
    )
    args = parser.parse_args()
    result = asyncio.run(run_acceptance(Path(args.data_dir)))
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

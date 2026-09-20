import csv
from datetime import UTC, datetime

from openpyxl import load_workbook

from app.domain.family import FamilyType, PatentFamily, PatentPublication
from app.library.export import EXPORT_COLUMNS
from app.library.service import PatentLibraryService
from app.library.store import SQLitePatentLibrary


def _service(tmp_path):
    store = SQLitePatentLibrary(tmp_path / "library.db")
    family = PatentFamily(
        family_type=FamilyType.DOCDB_SIMPLE,
        source="TEST",
        source_family_id="F1",
        members=[
            PatentPublication(
                publication_number="JP2024000123A",
                jurisdiction="JP",
                kind_code="A",
                title="Damper test",
                original_assignees=("Company A",),
            )
        ],
    )
    store.upsert_family(
        family,
        seen_at=datetime(2026, 9, 20, 8, 0, tzinfo=UTC),
    )
    store.add_tag("JP2024000123A", "重点")
    return PatentLibraryService(store), store


def test_csv_export_uses_expected_columns_and_utf8_bom(tmp_path):
    service, store = _service(tmp_path)
    path = service.export(tmp_path / "patents.csv")

    raw = path.read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf")

    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.DictReader(stream))
    assert tuple(rows[0].keys()) == EXPORT_COLUMNS
    assert rows[0]["publication_number"] == "JP2024000123A"
    assert rows[0]["tags"] == "重点"
    store.close()


def test_xlsx_export_is_native_and_readable(tmp_path):
    service, store = _service(tmp_path)
    path = service.export(tmp_path / "patents.xlsx")

    workbook = load_workbook(path)
    sheet = workbook["Patent Library"]

    assert sheet.freeze_panes == "A2"
    assert sheet.cell(1, 1).value == "publication_number"
    assert sheet.cell(2, 1).value == "JP2024000123A"
    assert sheet.auto_filter.ref is not None
    assert sheet.cell(1, 1).font.bold is True
    store.close()
